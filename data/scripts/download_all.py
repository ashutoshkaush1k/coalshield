"""Stage D2: download every automatic source in data/sources.yaml and check the manual ones.

    data\\.venv\\Scripts\\python.exe data\\scripts\\download_all.py [--only S06,S08] [--force S08]

Safe to re-run at any time:
  - a file already on disk whose SHA-256 matches the manifest is skipped, not re-downloaded;
  - an interrupted download resumes from its .part file (HTTP Range) when the server allows;
  - manual folders are re-checked every run, so files dropped in later are verified (checksum,
    size) and recorded, and anything still missing is marked pending-manual.
--force <IDs> re-downloads those sources (e.g. to take a newer weekly MSHA snapshot).

All traffic goes through PoliteSession: robots.txt, at most 1 request per second per host,
retries with backoff. Every downloaded file is checked against its expected type (a PDF must
start with %PDF, a zip or xlsx with PK), so an HTML error page served under a .pdf name is
rejected instead of saved. Writes data/raw/<source folder>/README.md for each source, updates
sources.yaml, and prints the summary table. Exit code 1 if any automatic download failed.
"""

from __future__ import annotations

import argparse
import datetime as dt
import os
import sys
import zipfile
from pathlib import Path
from urllib.parse import unquote, urlsplit

import requests

from common import DATA, load_config, load_manifest, save_manifest, sha256_file, slug, utc_now
from polite_http import PoliteSession, RobotsDisallowed

MAGIC = {".pdf": b"%PDF", ".zip": b"PK", ".xlsx": b"PK"}
EXT_FOR_TYPE = {"application/zip": ".zip", "application/x-zip-compressed": ".zip",
                "application/pdf": ".pdf", "text/html": ".html"}
IGNORE_NAMES = {"README.md", ".gitkeep"}
# S02's licence is recomputed from the workbook on every run; with no workbook present it goes back
# to this, so a file that is later removed can never leave a "confirmed" licence behind.
GEM_LICENCE_UNCONFIRMED = "CC BY 4.0 (to be confirmed against the notice inside the downloaded file)"


# --- naming -------------------------------------------------------------------------------
def filename_for(f: dict) -> str:
    """Stable local name: the URL's own file name when it has one, else a slug of the entry name."""
    if f.get("filename"):
        return f["filename"]
    parts = urlsplit(f["url"])
    base = unquote(Path(parts.path).name)
    ext = Path(base).suffix.lower()
    if base and ext and ext not in {".php", ".asp", ".aspx", ".htm", ".html"} and not parts.query:
        return base
    ctype = (f.get("content_type") or "").split(";")[0].strip()
    return slug(f["name"]) + EXT_FOR_TYPE.get(ctype, ".html" if f.get("kind") == "page" else ".bin")


def rel(path: Path) -> str:
    return path.relative_to(DATA).as_posix()


# --- one automatic file -------------------------------------------------------------------
def fetch(session: PoliteSession, f: dict, force: bool) -> str:
    dest = DATA / f["target"] / filename_for(f)
    dest.parent.mkdir(parents=True, exist_ok=True)
    f["local_path"] = rel(dest)

    if dest.exists() and not force:
        recorded = f.get("sha256")
        actual = sha256_file(dest)
        if recorded and recorded == actual:
            f["download_status"] = "ok"
            return "present"
        if not recorded and f.get("size_bytes") and dest.stat().st_size == f["size_bytes"]:
            # Completed on an earlier run that ended before the manifest was saved.
            _record(f, dest, actual, None)
            return "present"

    part = dest.with_name(dest.name + ".part")
    offset = part.stat().st_size if part.exists() and not force else 0
    if force and part.exists():
        part.unlink()
    # identity: bytes on the wire are the bytes on disk. With gzip, a resume offset (counted in
    # decoded bytes) would not match the server's Range (counted in encoded bytes).
    headers = {"Accept-Encoding": "identity"}
    if offset:
        headers["Range"] = f"bytes={offset}-"
    resp = session.request("GET", f["url"], headers=headers, stream=True, timeout=120)
    try:
        if resp.status_code == 206 and offset:
            mode = "ab"                      # server honoured the resume
        elif resp.status_code == 200:
            mode, offset = "wb", 0           # full body (no resume support, or a fresh start)
        else:
            raise RuntimeError(f"HTTP {resp.status_code}")
        with part.open(mode) as fh:
            for chunk in resp.iter_content(chunk_size=1 << 20):
                fh.write(chunk)
    finally:
        resp.close()

    expected = MAGIC.get(dest.suffix.lower())
    with part.open("rb") as fh:
        head = fh.read(8)
    if expected and not head.startswith(expected):
        part.unlink()
        raise RuntimeError(f"server returned {head[:8]!r}..., not a {dest.suffix} file (error page?)")
    os.replace(part, dest)
    _record(f, dest, sha256_file(dest), resp.headers.get("Last-Modified"))
    return "downloaded"


def _record(f: dict, dest: Path, sha: str, last_modified: str | None) -> None:
    probed = f.get("size_bytes")
    f.update({"local_path": rel(dest), "sha256": sha, "size_bytes": dest.stat().st_size,
              "downloaded_at": utc_now(), "download_status": "ok"})
    if last_modified:
        f["last_modified"] = last_modified
    if probed and probed != f["size_bytes"]:
        f["size_note"] = f"server announced {probed} bytes at probe time; downloaded {f['size_bytes']}"
    f.pop("download_error", None)


# --- manual files -------------------------------------------------------------------------
def check_manual(f: dict, automatic_paths: set[str]) -> list[Path]:
    """Files present for a manual entry: an exact file, or anything new in its target folder."""
    target = DATA / f["target"]
    if target.suffix and not f["target"].endswith("/"):
        found = [target] if target.exists() else []
    else:
        found = sorted(p for p in target.rglob("*") if p.is_file() and p.name not in IGNORE_NAMES
                       and not p.name.endswith(".part") and rel(p) not in automatic_paths) if target.exists() else []
    previous = {x["path"]: x for x in f.get("found_files") or []}
    records = []
    for p in found:
        sha = sha256_file(p)
        old = previous.get(rel(p))
        records.append({"path": rel(p), "sha256": sha, "size_bytes": p.stat().st_size,
                        "recorded_at": old["recorded_at"] if old and old["sha256"] == sha else utc_now()})
    if records:
        f["found_files"] = records
        f["download_status"] = "present-manual"
    else:
        f.pop("found_files", None)
        f["download_status"] = "pending-manual"
    return found


def gem_licence(xlsx: Path) -> dict:
    """Read the licence notice inside the GEM workbook (S02) - the source of truth for its licence."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    hits = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and any(k in v.lower() for k in ("licen", "creative commons", "cc by", "attribution")):
                    hits.append({"sheet": ws.title, "cell": cell.coordinate, "text": " ".join(v.split())[:400]})
            if len(hits) >= 10:
                break
    wb.close()
    return {"checked_file": rel(xlsx), "notices": hits}


def ensure_manual_folders(manifest: dict) -> list[str]:
    """Create every raw/ folder MANUAL_STEPS.md tells a person to save into, each with a .gitkeep,
    so nobody ever has to create a folder by hand (and the folders exist in a fresh clone).

    Paths are read from MANUAL_STEPS.md itself, then checked against the manifest's manual targets;
    a folder named in one but not the other is reported, because the two should never drift apart.
    """
    import re
    text = (DATA / "MANUAL_STEPS.md").read_text(encoding="utf-8")
    from_doc = {m.replace("\\", "/").rstrip("/") for m in re.findall(r"data\\(raw(?:\\[\w.\-]+)+)", text)}
    from_manifest = set()
    for src in manifest["sources"]:
        for f in src.get("files") or []:
            if f.get("manual") or (src.get("manual") and not f.get("url")):
                t = f["target"].rstrip("/")
                from_manifest.add(t if f["target"].endswith("/") else str(Path(t).parent).replace("\\", "/"))
    for only_doc in sorted(from_doc - from_manifest):
        print(f"  note: {only_doc} is in MANUAL_STEPS.md but not a manual target in sources.yaml")
    for only_manifest in sorted(from_manifest - from_doc):
        print(f"  note: {only_manifest} is a manual target in sources.yaml but not in MANUAL_STEPS.md")
    created = []
    for folder in sorted(from_doc | from_manifest):
        path = DATA / folder
        path.mkdir(parents=True, exist_ok=True)
        keep = path / ".gitkeep"
        if not keep.exists():
            keep.write_text("", encoding="utf-8")
            created.append(folder)
    return created


# --- post-processing ----------------------------------------------------------------------
def extract_ppe(zip_path: Path) -> dict:
    """S13: keep only dataset/ from the repository archive (the video and weights are not needed)."""
    out = zip_path.parent / "dataset"
    marker = out / ".extracted_from"
    sha = sha256_file(zip_path)
    if marker.exists() and marker.read_text(encoding="utf-8").strip() == sha:
        return {"extracted_to": rel(out), "status": "already extracted"}
    count = 0
    with zipfile.ZipFile(zip_path) as zf:
        for info in zf.infolist():
            parts = Path(info.filename).parts
            if len(parts) > 2 and parts[1] == "dataset" and not info.is_dir():
                dest = out.joinpath(*parts[2:])
                dest.parent.mkdir(parents=True, exist_ok=True)
                with zf.open(info) as src, dest.open("wb") as dst:
                    dst.write(src.read())
                count += 1
    marker.write_text(sha + "\n", encoding="utf-8")
    return {"extracted_to": rel(out), "files": count, "status": "extracted"}


def msha_snapshot(files: list[dict]) -> dict:
    """Which weekly MSHA release was taken. MSHA re-publishes every Friday afternoon, so later
    stages are reproducible only against this snapshot; it is identified by the servers'
    Last-Modified dates of the zips actually downloaded."""
    from email.utils import parsedate_to_datetime
    snap, weeks = {}, set()
    for f in files:
        if f.get("local_path", "").endswith(".zip") and f.get("last_modified"):
            when = parsedate_to_datetime(f["last_modified"])
            year, week, _ = when.isocalendar()
            weeks.add(f"{year}-W{week:02d}")
            snap[f["name"]] = (f"published {when.astimezone(dt.timezone.utc).isoformat()} (server Last-Modified, "
                               f"ISO week {year}-W{week:02d}), downloaded {f.get('downloaded_at', '-')}")
    if weeks:
        snap["release_week"] = ", ".join(sorted(weeks)) + (" - files come from different weekly releases" if len(weeks) > 1 else "")
    snap["refresh_rule"] = "MSHA updates every Friday afternoon; re-run with --force S08 to take a newer release"
    return snap


# --- README per source --------------------------------------------------------------------
def write_readme(src: dict) -> None:
    folder = DATA / src["target_path"]
    folder.mkdir(parents=True, exist_ok=True)
    lines = [f"# {src['id']} - {src['name']}", "",
             "Generated by `data/scripts/download_all.py` from `data/sources.yaml`. The files in this",
             "folder are not committed (gitignored); only this README is. Re-create them with",
             "`data\\run_data.bat download`.", "",
             f"- **What:** {src['purpose']}",
             f"- **Publisher:** {src['publisher']}",
             f"- **From:** {src['landing_page']}",
             f"- **Licence:** {src['licence']}" + (f" ({src['licence_source']})" if src.get("licence_source") else ""),
             f"- **Redistribution:** {src['redistribution']}",
             f"- **Status:** {src.get('d2_status', '-')}", ""]
    if src.get("snapshot"):
        lines += ["## Snapshot", ""] + [f"- {k}: {v}" for k, v in src["snapshot"].items()] + [""]
    if src.get("d2_detail"):
        lines += ["## Details", ""] + [f"- {k}: {v}" for k, v in src["d2_detail"].items()] + [""]
    lines += ["## Files", "", "| File | Status | Size (bytes) | Downloaded (UTC) | SHA-256 | Source URL |",
              "|---|---|---|---|---|---|"]
    for f in src.get("files") or []:
        if f.get("found_files"):
            for x in f["found_files"]:
                lines.append(f"| `{x['path']}` | present-manual | {x['size_bytes']} | {x['recorded_at']} | `{x['sha256']}` | added by hand |")
        elif f.get("local_path") and f.get("sha256"):
            lines.append(f"| `{f['local_path']}` | {f.get('download_status')} | {f.get('size_bytes')} | "
                         f"{f.get('downloaded_at', '-')} | `{f['sha256']}` | {f.get('final_url') or f.get('url')} |")
        else:
            lines.append(f"| {f['name']} | {f.get('download_status', '-')} | - | - | - | "
                         f"{f.get('url') or f.get('manual_hint') or 'see MANUAL_STEPS.md'} |")
    (folder / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")


# --- main ---------------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", default="", help="comma-separated source IDs")
    ap.add_argument("--force", default="", help="comma-separated source IDs to re-download")
    args = ap.parse_args()
    only = {s for s in args.only.split(",") if s}
    force = {s for s in args.force.split(",") if s}

    manifest, header = load_manifest()
    config = load_config()
    for folder in ensure_manual_folders(manifest):
        print(f"  created folder {folder}/ (with .gitkeep) for a manual step")
    session = PoliteSession(timeout=120)
    automatic_paths = set()
    summary = []

    for src in manifest["sources"]:
        if only and src["id"] not in only:
            continue
        files = src.get("files") or []
        failures, detail = 0, {}
        src.pop("d2_detail", None)   # rebuilt below from this run only - never carried over stale

        if src["access"] == "repo":
            src["d2_status"] = "ok (in repo)"
            summary.append((src["id"], src["name"], src["d2_status"], len(files), 0))
            continue
        if src["access"] == "skipped":
            src["d2_status"] = "skipped (see notes for fallback)"
            write_readme(src)
            summary.append((src["id"], src["name"], src["d2_status"], 0, 0))
            continue

        for f in files:
            if f.get("url") and not f.get("manual") and f.get("kind") != "api":
                try:
                    what = fetch(session, f, src["id"] in force)
                    automatic_paths.add(f["local_path"])
                    print(f"  {src['id']}  {what:10}  {f['size_bytes']:>12,}  {f['local_path']}", flush=True)
                except (requests.RequestException, RobotsDisallowed, RuntimeError, OSError) as exc:
                    failures += 1
                    f["download_status"] = "failed"
                    f["download_error"] = f"{type(exc).__name__}: {exc}"[:300]
                    print(f"  {src['id']}  FAILED      {f['name']}: {f['download_error']}", flush=True)
        # manual entries are checked after the automatic ones, so their folders can be told apart
        pending = present = 0
        for f in files:
            if f.get("manual") or (not f.get("url") and f.get("kind") != "api"):
                if not f.get("manual") and not src.get("manual"):
                    f["download_status"] = "not located"
                    continue
                if f.get("skipped"):
                    f["download_status"] = "skipped"      # decided by the user; see skip_reason
                    continue
                found = check_manual(f, automatic_paths)
                present += bool(found)
                pending += not found
                if src["id"] == "S02":
                    src["licence"] = GEM_LICENCE_UNCONFIRMED
                if found and src["id"] == "S02":
                    books = [p for p in found if p.suffix.lower() == ".xlsx"]
                    checks = [gem_licence(p) for p in books]
                    detail["licence_check"] = checks
                    def cc4(c):
                        text = " ".join(n["text"] for n in c["notices"]).lower()
                        return "4.0" in text and ("creative commons" in text or "cc by" in text)
                    confirmed = [c["checked_file"] for c in checks if cc4(c)]
                    if confirmed and len(confirmed) == len(checks):
                        src["licence"] = ("CC BY 4.0 - confirmed from the licence notice inside each downloaded "
                                          f"workbook ({len(confirmed)} of {len(checks)})")
                    elif confirmed:
                        src["licence"] = (f"CC BY 4.0 confirmed in {len(confirmed)} of {len(checks)} workbooks "
                                          "(see d2_detail.licence_check for the others)")
                    elif any(c["notices"] for c in checks):
                        src["licence"] = "Licence notice found in the workbooks but it is not CC BY 4.0 - review d2_detail"
                    else:
                        src["licence"] = "No licence notice found in the workbooks - still unconfirmed (CC BY 4.0 per GEM site)"
            elif f.get("kind") == "api" and src["id"] == "S09":
                try:
                    from download_openaq import run as openaq_run
                    result = openaq_run(config)
                except Exception as exc:  # the key is never part of an exception message
                    result = {"download_status": "failed", "message": f"{type(exc).__name__}: {str(exc)[:200]}"}
                    failures += 1
                f["download_status"] = result["download_status"]
                detail["openaq"] = result
                pending += result["download_status"] == "skipped-manual"
                print(f"  {src['id']}  {result['download_status']:10}  {result['message']}", flush=True)

        if src["id"] == "S13":
            z = next((DATA / f["local_path"] for f in files if f.get("local_path") and f.get("download_status") == "ok"), None)
            if z and z.suffix == ".zip":
                detail["extraction"] = extract_ppe(z)
        if src["id"] == "S08":
            src["snapshot"] = msha_snapshot(files)
            if not failures:
                try:
                    from filter_msha import run as msha_run
                    detail["filtered"] = msha_run(src, config)
                except ImportError:
                    detail["filtered"] = "filter_msha.py not present yet"

        auto_ok = [f for f in files if f.get("download_status") == "ok" and f.get("local_path")]
        n_files = len(auto_ok) + sum(len(f.get("found_files") or []) for f in files)
        size = sum(f.get("size_bytes") or 0 for f in auto_ok) + \
            sum(x["size_bytes"] for f in files for x in f.get("found_files") or [])
        if src["id"] == "S09" and detail.get("openaq", {}).get("download_status") == "ok":
            api_files = [p for p in (DATA / src["target_path"]).rglob("*.json")]
            n_files += len(api_files)
            size += sum(p.stat().st_size for p in api_files)
            detail["openaq"]["files_on_disk"] = len(api_files)
        if failures:
            status = "failed"
        elif not auto_ok and pending and not present:
            status = "skipped-manual"
        elif pending:
            status = f"ok (+{pending} pending-manual)"
        else:
            status = "ok"
        src["d2_status"] = status
        if detail:
            src["d2_detail"] = detail
        write_readme(src)
        summary.append((src["id"], src["name"], status, n_files, size))
        save_manifest(manifest, header)   # after every source, so an interruption loses nothing

    save_manifest(manifest, header)
    print("\n| Source | Status | Files | Size |\n|---|---|---|---|")
    for sid, name, status, n, size in summary:
        print(f"| {sid} {name[:48]} | {status} | {n} | {size / 1_048_576:.1f} MB |")
    return 1 if any(s[2] == "failed" for s in summary) else 0


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).parent))
    sys.exit(main())
