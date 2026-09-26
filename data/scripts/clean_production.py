"""Stage D3: reference/production_company_monthly.csv and reference/production_company_annual.csv.

Monthly (S06, Ministry of Coal "Monthly Coal Statistics", provisional)
  Page 1 "Coal Production" and the page headed "Lignite Production" (page 5 or 6) of each
  downloaded monthly PDF. Each
  PDF gives the report month and the same month a year earlier, so the four PDFs (Jun, Jul, Aug
  2026 and Sep 2025) give eight months: Jun-Aug 2025, Jun-Aug 2026, Sep 2024 and Sep 2025.
  September 2026 is not yet published; September 2025 is its seasonal stand-in (config.yaml
  window 2026-06-28 to 2026-09-25) and must be labelled as such by the generator.
  Company-level overburden is not published monthly (only for the top 35 mines, page 12), so
  the monthly file has no OB column.

Annual (S06, Coal Directory of India 2024-25, cdchap2.xlsx)
  Table 3.11 (sheet PT11) production, Table 3.20 (PT20) open cast / underground split, Table 3.22
  (PT22) overburden removal and stripping ratio - for 2022-23 to 2024-25 where published. This is
  where company overburden is published.
  Also, for the D4 generators: Table 3.7 (PT7) month-wise 2024-25 production (the seasonal profile
  for months with no published monthly figure) -> production_profile_2024_25.csv, and Table 3.24
  (PT24) output per manshift by company and mine type -> oms_company.csv.

Parsing (page 1 rows have missing cells when a value is zero)
  Every growth figure is printed after an arrow (up or down triangle); those pairs are removed
  first. The remaining numbers are read by count: 6 = target, month, achievement %, month a year
  earlier, year-to-date, year-to-date a year earlier; 5 = the same without the %; 4 = without
  target and %. Checks, all enforced:
    - achievement % = month / target (within rounding);
    - ECL+BCCL+CCL+NCL+WCL+SECL+MCL+NEC = the CIL row, for both months;
    - year-to-date of one report minus the previous report's = that month (Jul and Aug 2026).
"""

from __future__ import annotations

import csv
import re
import sys

import openpyxl
import pdfplumber

from common import DATA, sha256_file

MONTHLY = ["raw/moc/monthly/msg-setp25.pdf", "raw/moc/monthly/msg-june26.pdf",
           "raw/moc/monthly/msg-july26.pdf", "raw/moc/monthly/msg-Aug26.pdf"]
DIRECTORY = "raw/moc/coal_directory_2024_25/cdchap2.xlsx"
OUT_M = DATA / "reference/production_company_monthly.csv"
OUT_A = DATA / "reference/production_company_annual.csv"
OUT_P = DATA / "reference/production_profile_2024_25.csv"
OUT_O = DATA / "reference/oms_company.csv"
MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9,
          "oct": 10, "nov": 11, "dec": 12}
CIL_ROWS = ["ECL", "BCCL", "CCL", "NCL", "WCL", "SECL", "MCL", "NEC"]
# published row label -> (company_id, scope)
LABELS = {**{k: (k, "company") for k in CIL_ROWS if k != "NEC"},
          "NEC": ("CIL", "company (NEC row: CIL's own coalfield, no subsidiary)"),
          "CIL": ("CIL", "group_total (sum of the eight rows above)"),
          "SCCL": ("SCCL", "company"), "Captives/Others": ("", "captives and others (not our companies)"),
          "NLCIL": ("NLC", "company"), "Total CIL": ("CIL", "group_total"), "NLCIL(coal)": ("NLC", "company")}
M_COLUMNS = ["company_id", "row_label", "scope", "mineral", "year", "month", "production_mt", "target_mt",
             "status", "source_file", "source_page", "source_table", "source_column"]
A_COLUMNS = ["company_id", "row_label", "scope", "fiscal_year", "coal_mt", "coking_mt", "opencast_mt",
             "underground_mt", "ob_removal", "ob_production_opencast_mt", "stripping_ratio", "source_file",
             "source_tables"]


def report_month(text: str) -> tuple[int, int]:
    m = re.search(r"Monthly Coal Statistics for ([A-Za-z]+)'?\s*(\d{4})", text)
    if not m:
        raise ValueError("report month not found on page 1")
    return int(m.group(2)), MONTHS[m.group(1)[:3].lower()]


def parse_rows(text: str, labels: list[str]) -> dict[str, dict]:
    out = {}
    for line in text.splitlines():
        m = re.match(r"^(?:\d+\s+)?(" + "|".join(re.escape(l) for l in labels) + r")\s+(.*)$", line.strip())
        if not m or m.group(1) in out:
            continue
        rest = re.sub(r"[▲▼]\s*-?\d+(?:\.\d+)?", " ", m.group(2))   # drop growth figures
        nums = [float(x) for x in re.findall(r"-?\d+(?:\.\d+)?", rest)]
        if not nums:
            continue   # a chart legend line ("NLCIL GMDCL GIPCL ..."), not a table row
        keys ={6: ["target", "month", "pct", "prev", "ytd", "ytd_prev"], 5: ["target", "month", "prev", "ytd", "ytd_prev"],
                4: ["month", "prev", "ytd", "ytd_prev"]}.get(len(nums))
        if not keys:
            raise ValueError(f"cannot read row '{line.strip()}' ({len(nums)} numbers after removing growth)")
        row = dict(zip(keys, nums))
        # the month is printed to 2 decimals, so allow its rounding (0.005 Mt) on top of the % rounding
        tol = 0.6 + (0.005 / row["target"] * 100 if row.get("target") else 0)
        if "pct" in row and row["target"] and abs(row["month"] / row["target"] * 100 - row["pct"]) > tol:
            raise ValueError(f"achievement % does not match month/target in '{line.strip()}'")
        out[m.group(1)] = row
    return out


def monthly() -> list[dict]:
    rows, ytd = [], {}
    for rel in MONTHLY:
        pdf = pdfplumber.open(DATA / rel)
        p1 = pdf.pages[0].extract_text()
        lig_page = next((i + 1 for i in range(min(10, len(pdf.pages)))
                         if re.search(r"^Lignite Production\s*$", pdf.pages[i].extract_text() or "", re.M)), None)
        if lig_page is None:
            raise ValueError(f"{rel}: no page headed 'Lignite Production' in the first 10 pages")
        p5 = pdf.pages[lig_page - 1].extract_text()
        year, month = report_month(p1)
        mname = [k for k, v in MONTHS.items() if v == month][0].capitalize()
        coal = parse_rows(p1, CIL_ROWS + ["CIL", "SCCL", "Captives/Others"])
        missing = [l for l in CIL_ROWS + ["CIL", "SCCL"] if l not in coal]
        if missing:
            raise ValueError(f"{rel}: rows {missing} not found on page 1")
        for key in ("month", "prev"):
            s = round(sum(coal[l][key] for l in CIL_ROWS), 2)
            if abs(s - coal["CIL"][key]) > 0.03:
                raise ValueError(f"{rel}: CIL rows sum to {s}, CIL row says {coal['CIL'][key]} ({key})")
        lig = parse_rows(p5, ["NLCIL"])
        if "NLCIL" not in lig:
            raise ValueError(f"{rel}: NLCIL row not found on page {lig_page} (Lignite Production)")
        ytd[(year, month)] = {l: coal[l]["ytd"] for l in coal}
        sources = [((x, coal[x]), "coal", 1, "Coal Production") for x in coal]
        sources.append((("NLCIL", lig["NLCIL"]), "lignite", lig_page, "Lignite Production"))
        for (label, r), mineral, page, table in sources:
            cid, scope = LABELS[label]
            fy = year + 1 if month >= 4 else year
            for y, key, col in ((year, "month", f"Production during {mname} FY {str(fy)[2:]}"),
                                (year - 1, "prev", f"Production during {mname} FY {str(fy - 1)[2:]}")):
                rows.append({"company_id": cid, "row_label": label, "scope": scope, "mineral": mineral, "year": y,
                             "month": month, "production_mt": f"{r[key]:.2f}",
                             "target_mt": f"{r['target']:.2f}" if key == "month" and "target" in r else "",
                             "status": "provisional" if key == "month" else "as reported a year later",
                             "source_file": rel, "source_page": page, "source_table": table, "source_column": col})
    # year-to-date consistency across consecutive reports
    for (y, m), prev in (((2026, 7), (2026, 6)), ((2026, 8), (2026, 7))):
        for label in CIL_ROWS + ["CIL", "SCCL"]:
            got = round(ytd[(y, m)][label] - ytd[prev][label], 2)
            pub = next(float(r["production_mt"]) for r in rows if r["row_label"] == label and r["year"] == y
                       and r["month"] == m and r["mineral"] == "coal")
            if abs(got - pub) > 0.05:
                raise ValueError(f"{label} {y}-{m:02d}: YTD difference {got} != month {pub}")
    return rows


def cells(ws, first_data_row: int) -> list[list]:
    return [list(r) for r in ws.iter_rows(min_row=first_data_row, values_only=True)]


def annual() -> list[dict]:
    wb = openpyxl.load_workbook(DATA / DIRECTORY, read_only=True, data_only=True)
    years = ["2022-23", "2023-24", "2024-25"]
    want = {"ECL", "BCCL", "CCL", "NCL", "WCL", "SECL", "MCL", "NEC", "CIL", "Total CIL", "SCCL", "NLCIL"}
    data: dict[tuple[str, str], dict] = {}

    def put(label, fy, **kw):
        label = "CIL" if label == "Total CIL" else label
        d = data.setdefault((label, fy), {})
        d.update({k: v for k, v in kw.items() if v is not None})

    for r in cells(wb["PT11"], 6):   # Company | (coking, non-coking, total) x 3 years
        if r[0] in want:
            for i, fy in enumerate(years):
                put(r[0], fy, coking_mt=r[1 + 3 * i], coal_mt=r[3 + 3 * i])
    for r in cells(wb["PT20"], 7):   # Company | 2023-24 (OC qty, share, growth, UG qty, share, growth) | 2024-25 (...)
        if r[0] in want:
            for i, fy in enumerate(years[1:]):
                put(r[0], fy, opencast_mt=r[1 + 6 * i], underground_mt=r[4 + 6 * i])
    for r in cells(wb["PT22"], 6):   # Company | (OBR, OC production, stripping ratio) x 3 years
        if r[0] in want:
            for i, fy in enumerate(years):
                put(r[0], fy, ob_removal=r[1 + 3 * i], ob_production_opencast_mt=r[2 + 3 * i],
                    stripping_ratio=r[3 + 3 * i])
    rows = []
    for (label, fy), d in sorted(data.items(), key=lambda x: (x[0][1], x[0][0])):
        cid, scope = LABELS.get(label, (label, "company"))
        rows.append({"company_id": cid, "row_label": label, "scope": scope, "fiscal_year": fy,
                     **{k: ("" if d.get(k) in (None, "") else f"{float(d[k]):.3f}" if k != "stripping_ratio"
                            else f"{float(d[k]):.2f}")
                        for k in ["coal_mt", "coking_mt", "opencast_mt", "underground_mt", "ob_removal",
                                  "ob_production_opencast_mt", "stripping_ratio"]},
                     "source_file": DIRECTORY,
                     "source_tables": "Table 3.11 (PT11); Table 3.20 (PT20, 2023-24 and 2024-25 only); Table 3.22 (PT22)"})
    return rows


def profile() -> list[dict]:
    """Coal Directory Table 3.7 (PT7): month-wise raw coal production in 2024-25, CIL / SCCL / All India.
    Used as the seasonal profile for months with no published monthly figure (D4 full preset)."""
    wb = openpyxl.load_workbook(DATA / DIRECTORY, read_only=True, data_only=True)
    names = ["April", "May", "June", "July", "August", "September", "October", "November", "December",
             "January", "February", "March"]
    rows = []
    for r in cells(wb["PT7"], 6):
        if r[0] in names:
            m = names.index(r[0])
            rows.append({"fiscal_year": "2024-25", "month": (m + 3) % 12 + 1, "year": 2024 if m < 9 else 2025,
                         "cil_mt": f"{float(r[1]):.3f}", "sccl_mt": f"{float(r[3]):.3f}",
                         "all_india_mt": f"{float(r[11]):.3f}", "source_file": DIRECTORY,
                         "source_table": f"Table 3.7 (PT7), row {r[0]}"})
    if len(rows) != 12:
        raise ValueError(f"PT7: expected 12 months, read {len(rows)}")
    return rows


def oms() -> list[dict]:
    """Coal Directory Table 3.24 (PT24): production, manshifts and output per manshift (OMS) by company
    and mine type (OC / UG / ALL), 2024-25 columns. Used for manpower in D4."""
    wb = openpyxl.load_workbook(DATA / DIRECTORY, read_only=True, data_only=True)
    rows = []
    for r in cells(wb["PT24"], 6):
        label = (r[0] or "").strip()
        if label in {"ECL", "BCCL", "CCL", "NCL", "WCL", "SECL", "MCL", "NEC", "CIL", "SCCL"} and r[1] in {"OC", "UG", "ALL"}:
            prod, ms, o = r[8], r[9], r[10]
            cid, scope = LABELS.get(label, (label, "company"))
            rows.append({"company_id": cid, "row_label": label, "scope": scope, "fiscal_year": "2024-25",
                         "mine_type": r[1], "production_mt": "" if prod is None else f"{float(prod):.3f}",
                         "manshifts_million": "" if ms is None else f"{float(ms):.4f}",
                         "oms_t": "" if o in (None, 0) else f"{float(o):.3f}", "source_file": DIRECTORY,
                         "source_table": "Table 3.24 (PT24), 2024-25 columns"})
    return rows


def main() -> int:
    for rel in MONTHLY + [DIRECTORY]:
        if not (DATA / rel).exists():
            print(f"ERROR: {rel} missing - run data\\run_data.bat download", file=sys.stderr)
            return 1
    try:
        m, a, pr, om = monthly(), annual(), profile(), oms()
    except ValueError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 1
    m.sort(key=lambda r: (r["mineral"], r["year"], r["month"], r["row_label"]))
    for path, cols, rows in ((OUT_M, M_COLUMNS, m), (OUT_A, A_COLUMNS, a), (OUT_P, list(pr[0]), pr), (OUT_O, list(om[0]), om)):
        with path.open("w", encoding="utf-8", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=cols, lineterminator="\n")
            w.writeheader()
            w.writerows(rows)
    months = sorted({(r["year"], r["month"]) for r in m})
    print(f"Wrote reference/production_company_monthly.csv: {len(m)} rows, months "
          f"{', '.join(f'{y}-{mo:02d}' for y, mo in months)}; checks passed (CIL sums, % vs target, YTD differences); "
          f"sha256 {sha256_file(OUT_M)[:16]}")
    print(f"Wrote reference/production_company_annual.csv: {len(a)} rows (2022-23 to 2024-25), sha256 {sha256_file(OUT_A)[:16]}")
    print(f"Wrote reference/production_profile_2024_25.csv: 12 months (Table 3.7); reference/oms_company.csv: "
          f"{len(om)} rows (Table 3.24)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
